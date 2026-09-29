#!/usr/bin/env python3
"""报价工作簿体检：打印 sheet 列表、表头、数据验证允许值、最后数据行。

动手改文件之前先跑这个，拿到 C 列允许值清单和真实的最后数据行。

用法:
    python3 inspect_workbook.py <目标文件.xlsx> [--sheet Sheet1]
"""

import argparse
import sys
from pathlib import Path


def parse_allowed(formula1: str):
    """把 '"a,b,c"' 形式的 formula1 解析为允许值列表。"""
    if not formula1:
        return []
    s = formula1.strip()
    if s.startswith('"') and s.endswith('"'):
        s = s[1:-1]
    return [x.strip() for x in s.split(",") if x.strip()]


def last_data_row(ws):
    """以 A 列序号为准找真实最后数据行，避开空白样式行和合计行。"""
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            last = r
    return last


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path", help="xlsx 文件路径")
    ap.add_argument("--sheet", help="只检查指定 sheet")
    args = ap.parse_args()

    path = Path(args.path).expanduser()
    if not path.exists():
        sys.exit(f"文件不存在: {path}")

    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl 不可用，请改用 references/excel-xml-fallback.md 中的 XML 方案")

    wb = openpyxl.load_workbook(path)
    sheets = [wb[args.sheet]] if args.sheet else wb.worksheets

    print(f"文件: {path}")
    print(f"全部 sheet: {wb.sheetnames}\n")

    for ws in sheets:
        print(f"=== sheet: {ws.title} ===")
        print(f"表头: {[c.value for c in ws[1]]}")

        widths = {k: round(v.width, 2) for k, v in ws.column_dimensions.items() if v.width}
        print(f"列宽: {widths}")

        dvs = list(ws.data_validations.dataValidation)
        if not dvs:
            print("数据验证: 无（新建 sheet 时记得同步！）")
        for dv in dvs:
            print(f"数据验证 {dv.sqref}:")
            for v in parse_allowed(dv.formula1):
                print(f"    - {v}")

        last = last_data_row(ws)
        print(f"最后数据行: {last}  下一可写行: {last + 1}")

        total = 0.0
        print("已有明细:")
        for r in range(2, last + 1):
            seq = ws.cell(row=r, column=1).value
            kind = ws.cell(row=r, column=3).value
            detail = ws.cell(row=r, column=4).value
            days = ws.cell(row=r, column=5).value
            if isinstance(days, (int, float)):
                total += days
            color = ws.cell(row=r, column=4).font.color
            red = "  [红]" if color and color.rgb == "FFFF0000" else ""
            print(f"  行{r}: {seq} | {kind} | {days} | {detail}{red}")
        print(f"人天合计（按 E 列实数累加）: {total}")

        for ref in ("G1", "G2", "H1", "H2"):
            v = ws[ref].value
            if v is not None:
                print(f"{ref} = {v!r}")
        print()


if __name__ == "__main__":
    main()
