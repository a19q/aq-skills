#!/usr/bin/env python3
"""从 JSON 追加报价行：自动映射类型术语、按规则填人天、增补项标红、序号连续。

用法:
    python3 append_rows.py <目标文件.xlsx> --rows rows.json [--sheet Sheet1] [--dry-run]

rows.json 格式:
    [
      {"feature": "FR-02", "kind": "UI开发",   "detail": "标记工具栏UI开发",       "highlight": false},
      {"feature": "FR-02", "kind": "逻辑开发", "detail": "标记清除功能交互逻辑开发", "highlight": true}
    ]

kind 只能是 UI开发 / 组件开发 / 逻辑开发 / 接口对接。
脚本会读取目标 sheet 的数据验证允许值并映射；映射不到时报错退出，需与用户确认。
"""

import argparse
import json
import sys
from copy import copy
from pathlib import Path

# 固定人天规则
DAYS = {"UI开发": 0.5, "组件开发": 0.5, "逻辑开发": 1.0, "接口对接": 0.5}

# 术语 -> 模板允许值的匹配关键词（按优先级）
# 注意 接口对接 用完整的「軟件接口對接」，避免误匹配「接口」或「硬件接口對接」
KIND_KEYWORDS = {
    "UI开发": ["頁面", "页面", "Page", "UI"],
    "组件开发": ["頁面", "页面", "Page", "UI"],
    "逻辑开发": ["邏輯", "逻辑", "Logic"],
    "接口对接": ["軟件接口對接", "软件接口对接", "軟件接口", "软件接口"],
}

RED = "FFFF0000"


def parse_allowed(formula1: str):
    if not formula1:
        return []
    s = formula1.strip()
    if s.startswith('"') and s.endswith('"'):
        s = s[1:-1]
    return [x.strip() for x in s.split(",") if x.strip()]


def collect_allowed(ws):
    allowed = []
    for dv in ws.data_validations.dataValidation:
        allowed.extend(parse_allowed(dv.formula1))
    return allowed


def map_kind(kind, allowed):
    """把技能术语映射为模板允许值；无允许值时直接返回原术语。"""
    if kind not in DAYS:
        sys.exit(f"非法 kind: {kind!r}，只能是 {list(DAYS)}")
    if not allowed:
        return kind
    for kw in KIND_KEYWORDS[kind]:
        for a in allowed:
            if kw in a:
                return a
    sys.exit(
        f"未找到 {kind} 对应的模板允许值，请与用户确认映射。\n"
        f"模板允许值: {allowed}"
    )


def last_data_row(ws):
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            last = r
    return last


def copy_style_from(ws, src_row, dst_row, col):
    if src_row < 2:
        return
    src = ws.cell(row=src_row, column=col)
    dst = ws.cell(row=dst_row, column=col)
    dst.font = copy(src.font)
    dst.alignment = copy(src.alignment)
    dst.border = copy(src.border)
    dst.number_format = src.number_format


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path", help="目标 xlsx 文件")
    ap.add_argument("--rows", required=True, help="rows.json 路径")
    ap.add_argument("--sheet", default=None, help="目标 sheet，默认第一个")
    ap.add_argument("--dry-run", action="store_true", help="只打印计划，不写文件")
    args = ap.parse_args()

    path = Path(args.path).expanduser()
    if not path.exists():
        sys.exit(f"文件不存在: {path}。先用 build_template.py 创建，或确认路径。")

    rows = json.loads(Path(args.rows).expanduser().read_text(encoding="utf-8"))
    if not isinstance(rows, list) or not rows:
        sys.exit("rows.json 必须是非空数组")

    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font
    except ImportError:
        sys.exit("openpyxl 不可用，请改用 references/excel-xml-fallback.md 中的 XML 方案")

    # 不使用 data_only=True，否则回写会抹掉合计公式
    wb = load_workbook(path)
    ws = wb[args.sheet] if args.sheet else wb.worksheets[0]

    allowed = collect_allowed(ws)
    if not allowed:
        print("警告: 目标 sheet 无数据验证，类型列将写入原术语", file=sys.stderr)

    last = last_data_row(ws)
    next_row = last + 1
    prev_seq = ws.cell(row=last, column=1).value if last > 1 else 0
    next_seq = (int(prev_seq) + 1) if isinstance(prev_seq, (int, float)) else 1

    plan = []
    for i, item in enumerate(rows):
        kind = item["kind"]
        plan.append({
            "row": next_row + i,
            "seq": next_seq + i,
            "feature": item.get("feature", ""),
            "kind_value": map_kind(kind, allowed),
            "detail": item["detail"],
            "days": DAYS[kind],
            "highlight": bool(item.get("highlight", False)),
        })

    added = sum(p["days"] for p in plan)
    print(f"目标: {path.resolve()}  sheet: {ws.title}")
    for p in plan:
        flag = " [红]" if p["highlight"] else ""
        print(f"  行{p['row']}: {p['seq']} | {p['kind_value']} | {p['days']} | {p['detail']}{flag}")
    print(f"本次新增人天: {added}")

    if args.dry_run:
        print("dry-run: 未写入")
        return

    for p in plan:
        r = p["row"]
        values = {1: p["seq"], 2: p["feature"], 3: p["kind_value"], 4: p["detail"], 5: p["days"]}
        for col, val in values.items():
            ws.cell(row=r, column=col, value=val)
            copy_style_from(ws, last, r, col)
        if p["highlight"]:
            d = ws.cell(row=r, column=4)
            f = d.font
            d.font = Font(name=f.name or "宋体", size=f.size or 11, bold=f.bold, color=RED)

    wb.save(path)

    # 复读校验
    wb2 = load_workbook(path)
    ws2 = wb2[ws.title]
    total = 0.0
    for r in range(2, last_data_row(ws2) + 1):
        v = ws2.cell(row=r, column=5).value
        if isinstance(v, (int, float)):
            total += v
    print(f"已写入 {len(plan)} 行 -> {path.resolve()}")
    print(f"当前总人天: {total}")


if __name__ == "__main__":
    main()
