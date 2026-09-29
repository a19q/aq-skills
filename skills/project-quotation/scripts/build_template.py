#!/usr/bin/env python3
"""重建等价的人天报价空模板。

用户未指定目录时，默认输出到当前系统的下载文件夹。

用法:
    python3 build_template.py                          # -> ~/Downloads/人天报价模板PMX.xlsx
    python3 build_template.py --out-dir .              # -> ./人天报价模板PMX.xlsx
    python3 build_template.py --name 某项目报价.xlsx --sheets Sheet1 Sheet2
"""

import argparse
import os
import sys
from pathlib import Path

COL_WIDTHS = {
    "A": 6.48076923076923,
    "B": 32.8557692307692,
    "C": 25.7980769230769,
    "D": 62.1826923076923,
    "E": 12.5,
    "F": 8.88461538461539,
    "G": 9.92307692307692,
}

HEADERS = {
    "A1": "序号",
    "B1": "功能",
    "C1": "类型",
    "D1": "报价明细",
    "E1": "人天",
    "G1": "Total",
}

ALLOWED = (
    "邏輯,軟件接口對接,頁面,定義一個實體（含CURD）,"
    "定義一個實體的DBUP,接口,硬件接口對接,硬件集成,Discount"
)


def default_download_dir() -> Path:
    """跨平台定位系统下载文件夹。"""
    if os.name == "nt":
        base = os.environ.get("USERPROFILE") or os.path.expanduser("~")
        return Path(base) / "Downloads"
    return Path.home() / "Downloads"


def build_sheet(ws, title):
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    ws.title = title

    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    font = Font(name="宋体", size=11, color="FF000000")
    align = Alignment(horizontal="center", vertical="center")

    for ref, value in HEADERS.items():
        c = ws[ref]
        c.value = value
        c.font = font
        c.alignment = align

    # 合计公式：不要在后续写入中覆盖它
    c = ws["G2"]
    c.value = "=SUM(E:E)"
    c.font = font
    c.alignment = align

    # 类型列数据验证（每个 sheet 都要单独添加，copy_worksheet 不会带过来）
    dv = DataValidation(type="list", formula1=f'"{ALLOWED}"', allow_blank=True)
    dv.sqref = "C2:C1048576"
    ws.add_data_validation(dv)

    # 行高保持默认，不显式设置


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir", default=None, help="输出目录，默认系统下载文件夹")
    ap.add_argument("--name", default="人天报价模板PMX.xlsx", help="文件名")
    ap.add_argument("--sheets", nargs="+", default=["Sheet1"], help="sheet 名称列表")
    ap.add_argument("--force", action="store_true", help="已存在时覆盖")
    args = ap.parse_args()

    try:
        from openpyxl import Workbook
    except ImportError:
        sys.exit("openpyxl 不可用，请改用 references/excel-xml-fallback.md 中的 XML 方案")

    out_dir = Path(args.out_dir).expanduser() if args.out_dir else default_download_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / args.name

    if out.exists() and not args.force:
        sys.exit(f"文件已存在，请在原文件上追加，或加 --force 覆盖: {out}")

    wb = Workbook()
    build_sheet(wb.active, args.sheets[0])
    for title in args.sheets[1:]:
        # 逐个新建而非 copy_worksheet，确保数据验证完整
        build_sheet(wb.create_sheet(), title)

    wb.save(out)
    print(f"saved: {out.resolve()}")
    print(f"sheets: {args.sheets}")


if __name__ == "__main__":
    main()
